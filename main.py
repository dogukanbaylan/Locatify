"""Google Maps business scraper with a small web UI and Excel export."""

from __future__ import annotations

import asyncio
import json
import logging
import random
import re
import unicodedata
import uuid
from pathlib import Path
from urllib.parse import parse_qs, quote_plus, unquote, urlparse

import pandas as pd
from fastapi import FastAPI, HTTPException, Query, Request
from fastapi.responses import FileResponse, StreamingResponse
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from playwright.async_api import TimeoutError as PlaywrightTimeout
from playwright.async_api import async_playwright
from pydantic import BaseModel, Field

BASE_DIR = Path(__file__).resolve().parent
EXPORT_DIR = BASE_DIR / "exports"
TEMPLATE_PATH = BASE_DIR / "templates" / "index.html"
EXPORT_DIR.mkdir(exist_ok=True)

EXCEL_COLUMNS = [
    "İşletme Adı",
    "Adres",
    "Telefon Numarası",
    "Web Sitesi",
    "Yıldız Puanı",
    "Yorum Sayısı",
    "Kategori",
    "Konum",
    "E-posta",
    "Place ID",
    "CID",
]

FILE_ID_RE = re.compile(r"[a-f0-9]{32}__[A-Za-z0-9._-]{1,80}\.(xlsx|csv)")
COORD_RE = re.compile(r"@(-?\d+(?:\.\d+)?),(-?\d+(?:\.\d+)?)")
DATA_COORD_RE = re.compile(r"!3d(-?\d+(?:\.\d+)?)!4d(-?\d+(?:\.\d+)?)")
CHIJ_RE = re.compile(r"(ChIJ[0-9A-Za-z_-]{10,})")
EMAIL_RE = re.compile(r"[A-Za-z0-9._%+-]+@[A-Za-z0-9.-]+\.[A-Za-z]{2,}")
USER_AGENTS = [
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/156.0.0.0 Safari/537.36",
    "Mozilla/5.0 (Macintosh; Intel Mac OS X 10_15_7) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/156.0.0.0 Safari/537.36",
    "Mozilla/5.0 (X11; Linux x86_64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/156.0.0.0 Safari/537.36",
]
END_OF_LIST_RE = re.compile(
    r"reached the end of the list|listenin sonuna ulaştınız|sonuna ulaştınız",
    re.IGNORECASE,
)
PLACE_ID_RE = re.compile(r"!1s(0x[0-9a-f]+:0x[0-9a-f]+)", re.IGNORECASE)

logging.basicConfig(
    level=logging.INFO,
    format="%(asctime)s %(levelname)s %(message)s",
)
log = logging.getLogger("maps-scraper")

app = FastAPI(title="Locatify", version="1.0.0")
scrape_lock = asyncio.Lock()


class SearchRequest(BaseModel):
    query: str = Field(..., min_length=2, max_length=200)
    headless: bool = True
    max_results: int = Field(15, ge=1, le=40)
    delay: float = Field(3, ge=0, le=60)
    export_format: str = "xlsx"
    proxy: str = ""


def validate_params(
    query: str,
    max_results: int,
    delay: float,
    export_format: str = "xlsx",
    proxy: str = "",
) -> tuple[str, float, str, str]:
    query = (query or "").strip()
    if len(query) < 2:
        raise HTTPException(status_code=400, detail="Enter a search of at least 2 characters.")
    if len(query) > 200:
        raise HTTPException(status_code=400, detail="Search text is too long (max 200 characters).")
    if not 1 <= max_results <= 40:
        raise HTTPException(status_code=400, detail="max_results must be between 1 and 40.")
    if not 0 <= delay <= 60:
        raise HTTPException(status_code=400, detail="delay must be between 0 and 60 seconds.")
    export_format = (export_format or "xlsx").strip().lower()
    if export_format not in {"xlsx", "csv"}:
        raise HTTPException(status_code=400, detail="export_format must be xlsx or csv.")
    proxy = (proxy or "").strip()
    if len(proxy) > 300:
        raise HTTPException(status_code=400, detail="Proxy URL is too long.")
    if proxy:
        parse_proxy(proxy)
    return query, float(delay), export_format, proxy


def slugify(text: str) -> str:
    text = text.strip().replace("İ", "i").replace("ı", "i").lower()
    text = unicodedata.normalize("NFKD", text)
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    text = re.sub(r"[^a-z0-9]+", "-", text).strip("-")
    return (text[:50] or "maps-results")


def digits_only(value: str) -> str:
    return re.sub(r"\D", "", value or "")


def clean_lines(raw: str, labels: set[str]) -> list[str]:
    lines: list[str] = []
    for line in (raw or "").replace("\r", "").split("\n"):
        line = "".join(ch for ch in line if ch.isprintable()).strip(" ·\u00b7|-")
        if not line:
            continue
        lowered = line.lower().rstrip(":")
        if lowered in labels:
            continue
        line = re.sub(
            r"^(address|adres|phone|telefon|website|web sitesi)\s*:\s*",
            "",
            line,
            flags=re.IGNORECASE,
        ).strip()
        if line and line.lower() not in labels:
            lines.append(line)
    return list(dict.fromkeys(lines))


def pick_address(text: str, aria: str) -> str:
    labels = {"address", "adres", "directions", "yol tarifi", "copy address", "adresi kopyala"}
    for raw in (text, aria):
        lines = clean_lines(raw, labels)
        if lines:
            return ", ".join(lines)
    return ""


def pick_phone(text: str, item_id: str) -> str:
    labels = {"phone", "telefon", "call", "ara", "copy phone number", "telefon numarasını kopyala"}
    cleaned = ""
    lines = clean_lines(text, labels)
    if lines:
        cleaned = lines[-1]
    if len(digits_only(cleaned)) >= 7:
        return cleaned
    if item_id and "tel:" in item_id:
        raw = unquote(item_id.split("tel:", 1)[1]).strip()
        if len(digits_only(raw)) >= 7:
            return raw
    return ""


def pick_website(href: str, text: str) -> str:
    href = (href or "").strip()
    text = " ".join((text or "").split())
    if href.startswith("http") and "google." not in urlparse(href).netloc:
        return href
    if "google." in (urlparse(href).netloc if href.startswith("http") else ""):
        query = parse_qs(urlparse(href).query)
        target = (query.get("q") or query.get("url") or [""])[0]
        if target.startswith("http") and "google." not in urlparse(target).netloc:
            return target
    if text and " " not in text and "." in text and "google." not in text.lower():
        return text if text.startswith("http") else f"https://{text.lstrip('/')}"
    return ""


def pick_rating(visible: str, aria: str) -> str:
    visible = (visible or "").strip()
    if re.search(r"\d", visible) and len(visible) <= 6:
        return visible
    match = re.search(r"\d+(?:[.,]\d+)?", aria or "")
    return match.group(0) if match else ""


def pick_review_count(visible: str, aria: str) -> str:
    for raw in (aria, visible):
        if not raw:
            continue
        match = re.search(r"(\d[\d.,\s]*)", raw)
        if match:
            return match.group(1).strip(" .,")
    return ""


def review_from_card(text: str) -> str:
    """Pull a review count from a result card when the place page does not show one."""
    lines = [line.strip() for line in (text or "").splitlines() if line.strip()]
    for index, line in enumerate(lines):
        if re.fullmatch(r"\d+[.,]\d+", line) and index + 1 < len(lines):
            nxt = lines[index + 1]
            if re.fullmatch(r"\(?\d[\d.,]*\)?", nxt):
                return nxt.strip("()")
        labeled = re.search(r"(\d[\d.,]*)\s+(?:reviews|yorum)\b", line, flags=re.IGNORECASE)
        if labeled:
            return labeled.group(1)
    return ""


def place_key(href: str) -> str:
    match = PLACE_ID_RE.search(href or "")
    if match:
        return match.group(1).lower()
    slug = re.search(r"/maps/place/([^/]+)", href or "")
    if slug:
        return unquote(slug.group(1)).lower()
    return (href or "").split("?")[0]


def blank_place(name: str = "") -> dict[str, str]:
    return {
        "name": name,
        "address": "",
        "phone": "",
        "website": "",
        "rating": "",
        "review_count": "",
        "category": "",
        "maps_url": "",
        "email": "",
        "place_id": "",
        "cid": "",
    }


def clean_email(value: str) -> str:
    match = EMAIL_RE.search(value or "")
    if not match:
        return ""
    email = match.group(0).strip(".").lower()
    if email.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg")):
        return ""
    return email


def coords_from_url(url: str) -> tuple[str, str]:
    decoded = unquote(url or "")
    match = COORD_RE.search(decoded) or DATA_COORD_RE.search(decoded)
    if not match:
        return "", ""
    return match.group(1), match.group(2)


def ids_from_url(url: str) -> tuple[str, str]:
    decoded = unquote(url or "")
    feature = PLACE_ID_RE.search(decoded)
    chij = CHIJ_RE.search(decoded)
    place_id = chij.group(1) if chij else (feature.group(1) if feature else "")
    cid = ""
    if feature:
        hex_cid = feature.group(1).split(":")[-1]
        try:
            cid = str(int(hex_cid, 16))
        except ValueError:
            cid = ""
    return place_id, cid


def maps_place_link(url: str) -> str:
    decoded = unquote(url or "").strip()
    if "/maps/" not in decoded:
        return ""
    parsed = urlparse(decoded if "://" in decoded else f"https://{decoded}")
    host = parsed.netloc or "www.google.com"
    if "google." not in host or not parsed.path:
        return ""
    return f"https://{host}{parsed.path}"


def apply_url_meta(place: dict[str, str], url: str) -> None:
    link = maps_place_link(url)
    if link and "/maps/place/" in link:
        place["maps_url"] = link
    elif link and not place.get("maps_url"):
        place["maps_url"] = link
    if not place.get("maps_url"):
        latitude, longitude = coords_from_url(url)
        if latitude and longitude:
            place["maps_url"] = f"https://www.google.com/maps/@{latitude},{longitude},17z"
    place_id, cid = ids_from_url(url)
    if place_id and not place.get("place_id"):
        place["place_id"] = place_id
    if cid and not place.get("cid"):
        place["cid"] = cid


def parse_proxy(proxy: str) -> dict[str, str]:
    raw = proxy.strip()
    if "://" not in raw:
        raw = "http://" + raw
    parsed = urlparse(raw)
    if parsed.scheme not in {"http", "https", "socks5"} or not parsed.hostname or not parsed.port:
        raise HTTPException(
            status_code=400,
            detail="Proxy URL must look like http://username:password@host:port",
        )
    config = {"server": f"{parsed.scheme}://{parsed.hostname}:{parsed.port}"}
    if parsed.username:
        config["username"] = unquote(parsed.username)
    if parsed.password:
        config["password"] = unquote(parsed.password)
    return config


def to_excel_row(place: dict[str, str]) -> dict[str, str]:
    return {
        "İşletme Adı": place.get("name", ""),
        "Adres": place.get("address", ""),
        "Telefon Numarası": place.get("phone", ""),
        "Web Sitesi": place.get("website", ""),
        "Yıldız Puanı": place.get("rating", ""),
        "Yorum Sayısı": place.get("review_count", ""),
        "Kategori": place.get("category", ""),
        "Konum": place.get("maps_url", ""),
        "E-posta": place.get("email", ""),
        "Place ID": place.get("place_id", ""),
        "CID": place.get("cid", ""),
    }


def save_export(rows: list[dict[str, str]], query: str, export_format: str) -> dict[str, str]:
    frame = pd.DataFrame([to_excel_row(row) for row in rows], columns=EXCEL_COLUMNS)
    frame = frame.fillna("").astype(str)
    slug = slugify(query)
    extension = "csv" if export_format == "csv" else "xlsx"
    file_id = f"{uuid.uuid4().hex}__{slug}.{extension}"
    path = EXPORT_DIR / file_id
    if extension == "csv":
        frame.to_csv(path, index=False, encoding="utf-8-sig")
    else:
        frame.to_excel(path, index=False, sheet_name="Sonuçlar")
        style_workbook(path)
    return {
        "file_id": file_id,
        "download_url": f"/api/download/{file_id}",
        "filename": f"{slug}.{extension}",
    }


def style_workbook(path: Path) -> None:
    workbook = load_workbook(path)
    sheet = workbook.active
    header_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF", name="Calibri")
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")
    widths = [36, 52, 22, 38, 16, 16, 28, 56, 32, 42, 22]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    link_font = Font(color="0F766E", underline="single", name="Calibri")
    body_font = Font(name="Calibri")
    for row in sheet.iter_rows(min_row=2, max_col=len(EXCEL_COLUMNS)):
        for cell in row:
            cell.number_format = "@"
            cell.font = body_font
            cell.alignment = Alignment(vertical="center", wrap_text=(cell.column in {1, 2}))
        for cell in (row[3], row[7]):
            value = str(cell.value or "")
            if value.startswith("http"):
                cell.hyperlink = value
                cell.font = link_font
    sheet.auto_filter.ref = sheet.dimensions
    sheet.freeze_panes = "A2"
    sheet.row_dimensions[1].height = 22
    workbook.save(path)


def resolve_export(file_id: str) -> Path:
    if not FILE_ID_RE.fullmatch(file_id):
        raise HTTPException(status_code=404, detail="File not found.")
    path = (EXPORT_DIR / file_id).resolve()
    if path.parent != EXPORT_DIR.resolve() or not path.is_file():
        raise HTTPException(status_code=404, detail="File not found.")
    return path


def friendly_error(exc: Exception) -> str:
    if isinstance(exc, PlaywrightTimeout):
        return (
            "Timed out waiting for Google Maps. Try again, or turn off headless mode "
            "to watch the browser. A consent dialog or bot check may be blocking the page."
        )
    message = str(exc).strip() or exc.__class__.__name__
    if "Executable doesn't exist" in message or "browserType.launch" in message:
        return "Playwright browser is missing. Run: python -m playwright install chromium"
    return message


EXTRACT_JS = """
() => {
  const textOf = (el) => (el ? (el.innerText || el.textContent || "") : "");
  const h1 = document.querySelector("h1");
  const name = textOf(h1).trim();

  let rating = "";
  const ratingNode = document.querySelector("div.F7nice span[aria-hidden='true']");
  if (ratingNode) rating = (ratingNode.textContent || "").trim();

  let ratingAria = "";
  const star = [...document.querySelectorAll("[role='img'][aria-label]")].find((el) =>
    /star|yıldız/i.test(el.getAttribute("aria-label") || "")
  );
  if (star) ratingAria = star.getAttribute("aria-label") || "";

  let reviewText = "";
  const nice = document.querySelector("div.F7nice");
  if (nice) {
    const hidden = [...nice.querySelectorAll("span[aria-hidden='true']")].map((el) =>
      (el.textContent || "").trim()
    );
    reviewText = hidden.find((value) => value && value !== rating) || "";
  }
  const reviewButton = [...document.querySelectorAll("button[aria-label]")].find((el) =>
    /review|yorum/i.test(el.getAttribute("aria-label") || "")
  );
  const reviewAria = reviewButton ? reviewButton.getAttribute("aria-label") || "" : "";

  const categoryButton =
    document.querySelector("button.DkEaL") ||
    document.querySelector("button[jsaction*='category']");
  const category = textOf(categoryButton).trim();

  const addressEl = document.querySelector("button[data-item-id='address']");
  const phoneEl = document.querySelector("button[data-item-id^='phone:tel:']");
  const webEl = document.querySelector("a[data-item-id='authority']");

  return {
    name,
    rating,
    ratingAria,
    reviewText,
    reviewAria,
    category,
    addressText: textOf(addressEl),
    addressAria: addressEl ? addressEl.getAttribute("aria-label") || "" : "",
    phoneText: textOf(phoneEl),
    phoneId: phoneEl ? phoneEl.getAttribute("data-item-id") || "" : "",
    websiteHref: webEl ? webEl.getAttribute("href") || "" : "",
    websiteText: textOf(webEl),
    email: (() => {
      const mail = document.querySelector("a[href^='mailto:']");
      if (mail) return (mail.getAttribute("href") || "").replace(/^mailto:/i, "").split("?")[0];
      const panel = document.querySelector("div[role='main']") || document.body;
      const found = (panel.innerText || "").match(/[A-Z0-9._%+-]+@[A-Z0-9.-]+\\.[A-Z]{2,}/i);
      return found ? found[0] : "";
    })(),
    href: location.href,
  };
}
"""

LINKS_JS = """
() => {
  const anchors = [...document.querySelectorAll("div[role='feed'] a[href*='/maps/place/']")];
  return anchors.map((anchor) => {
    const card = anchor.closest("[role='article']") || anchor.parentElement;
    return {
      href: anchor.href,
      name: anchor.getAttribute("aria-label") || "",
      cardText: card ? (card.innerText || "") : "",
    };
  });
}
"""


async def dismiss_consent(page) -> None:
    labels = (
        "Accept all",
        "Reject all",
        "I agree",
        "Tümünü kabul et",
        "Tümünü reddet",
        "Kabul et",
        "Kabul ediyorum",
    )
    frames = [page.main_frame, *page.frames]
    for frame in frames:
        for label in labels:
            button = frame.locator(f'button:has-text("{label}")')
            try:
                if await button.count():
                    await button.first.click(timeout=1500)
                    await page.wait_for_timeout(600)
                    return
            except Exception:
                continue


async def collect_links(page, max_results: int, progress, stop: asyncio.Event) -> tuple[list[dict[str, str]], bool]:
    """Return (links, single_place). single_place is true when Maps opened one listing directly."""
    await progress({"type": "status", "message": "Loading search results…"})
    try:
        await page.wait_for_selector("div[role='feed']", timeout=20000)
    except PlaywrightTimeout:
        if await page.locator("h1").count():
            return [], True
        title = await page.title()
        raise RuntimeError(
            f"Google Maps results did not appear (page: {title}). "
            "A consent screen or bot check may be in the way. Try again with the browser visible."
        )

    collected: dict[str, dict[str, str]] = {}
    stagnant = 0
    for _ in range(30):
        if stop.is_set():
            break
        batch = await page.evaluate(LINKS_JS)
        before = len(collected)
        for item in batch:
            href = item.get("href") or ""
            if "/maps/place/" not in href:
                continue
            collected.setdefault(
                place_key(href),
                {
                    "href": href,
                    "name": (item.get("name") or "").strip(),
                    "card_text": item.get("cardText") or "",
                },
            )
        await progress(
            {
                "type": "status",
                "message": f"Scrolling results… {min(len(collected), max_results)} listings found",
            }
        )
        if len(collected) >= max_results:
            break
        end_marker = page.get_by_text(END_OF_LIST_RE)
        if await end_marker.count():
            break
        if len(collected) == before:
            stagnant += 1
            if stagnant >= 3:
                break
        else:
            stagnant = 0
        await page.locator("div[role='feed']").evaluate(
            """(el) => {
                const links = el.querySelectorAll("a[href*='/maps/place/']");
                const last = links[links.length - 1];
                if (last) last.scrollIntoView({ block: "end" });
                el.scrollTop = el.scrollHeight;
            }"""
        )
        await page.wait_for_timeout(900)

    links = list(collected.values())[:max_results]
    return links, False


def normalize_place(raw: dict, fallback_name: str) -> dict[str, str]:
    name = " ".join((raw.get("name") or fallback_name or "").split())
    place = blank_place(name)
    place["address"] = pick_address(raw.get("addressText", ""), raw.get("addressAria", ""))
    place["phone"] = pick_phone(raw.get("phoneText", ""), raw.get("phoneId", ""))
    place["website"] = pick_website(raw.get("websiteHref", ""), raw.get("websiteText", ""))
    place["rating"] = pick_rating(raw.get("rating", ""), raw.get("ratingAria", ""))
    place["review_count"] = pick_review_count(raw.get("reviewText", ""), raw.get("reviewAria", ""))
    place["category"] = " ".join((raw.get("category") or "").split())
    place["email"] = clean_email(raw.get("email", ""))
    return place


async def read_open_place(page, fallback_name: str) -> dict[str, str]:
    await page.wait_for_selector("h1", timeout=15000)
    try:
        await page.wait_for_selector(
            "button[data-item-id='address'], button[data-item-id^='phone:tel:'], button.DkEaL",
            timeout=5000,
        )
    except PlaywrightTimeout:
        pass
    await page.wait_for_timeout(400)
    raw = await page.evaluate(EXTRACT_JS)
    place = normalize_place(raw, fallback_name)
    apply_url_meta(place, raw.get("href") or page.url)
    apply_url_meta(place, page.url)
    return place


def with_english_ui(url: str) -> str:
    if "hl=" in url:
        return url
    joiner = "&" if "?" in url else "?"
    return f"{url}{joiner}hl=en"


async def wait_between_listings(delay: float, stop: asyncio.Event) -> None:
    if stop.is_set():
        return
    await asyncio.sleep(random.uniform(delay, delay + 2))


async def scrape_maps(
    query: str,
    headless: bool,
    max_results: int,
    delay: float,
    proxy: str,
    progress,
    stop: asyncio.Event,
) -> list[dict[str, str]]:
    rows: list[dict[str, str]] = []
    await progress({"type": "status", "message": "Opening Google Maps…", "log": True})
    async with async_playwright() as playwright:
        launch_args: dict = {
            "headless": headless,
            "args": ["--disable-blink-features=AutomationControlled", "--lang=en-US"],
        }
        if proxy:
            launch_args["proxy"] = parse_proxy(proxy)
            log.info("Using a proxy for this search")
        browser = await playwright.chromium.launch(**launch_args)
        try:
            context = await browser.new_context(
                locale="en-US",
                user_agent=random.choice(USER_AGENTS),
                viewport={"width": 1440, "height": 900},
                extra_http_headers={"Accept-Language": "en-US,en;q=0.9"},
            )
            context.set_default_timeout(20000)
            page = await context.new_page()
            search_url = f"https://www.google.com/maps/search/{quote_plus(query)}?hl=en"
            await page.goto(search_url, wait_until="domcontentloaded", timeout=60000)
            await page.wait_for_timeout(800)
            await dismiss_consent(page)

            links, single_place = await collect_links(page, max_results, progress, stop)
            if stop.is_set():
                return rows

            if single_place:
                await wait_between_listings(delay, stop)
                if stop.is_set():
                    return rows
                place = await read_open_place(page, query)
                apply_url_meta(place, page.url)
                if place["name"]:
                    rows.append(place)
                    await progress(
                        {
                            "type": "place",
                            "index": 1,
                            "total": 1,
                            "place": place,
                            "message": f"Scraping 1/1: {place['name']}",
                            "log": True,
                        }
                    )
                return rows

            if not links:
                return rows

            total = len(links)
            await progress(
                {
                    "type": "status",
                    "message": f"Scraping 0/{total}: reading business details",
                    "total": total,
                    "log": True,
                }
            )
            for index, link in enumerate(links, start=1):
                if stop.is_set():
                    break
                await wait_between_listings(delay, stop)
                if stop.is_set():
                    break
                fallback_name = link.get("name") or ""
                status = f"Scraping {index}/{total}: {fallback_name or 'business'}"
                await progress(
                    {
                        "type": "status",
                        "message": status,
                        "index": index,
                        "total": total,
                        "log": True,
                    }
                )
                try:
                    await page.goto(with_english_ui(link["href"]), wait_until="domcontentloaded", timeout=45000)
                    await dismiss_consent(page)
                    place = await read_open_place(page, fallback_name)
                    apply_url_meta(place, link.get("href", ""))
                except Exception as exc:
                    log.warning("Skipped listing %s: %s", fallback_name, exc)
                    place = blank_place(fallback_name)
                    apply_url_meta(place, link.get("href", ""))
                if not place["name"]:
                    continue
                if not place["review_count"]:
                    place["review_count"] = review_from_card(link.get("card_text", ""))
                rows.append(place)
                log.info("Scraped %s (%s)", place["name"], place["phone"] or "no phone")
                await progress(
                    {
                        "type": "place",
                        "index": index,
                        "total": total,
                        "place": place,
                        "message": f"Scraped {index}/{total}: {place['name']}",
                        "log": True,
                    }
                )
            return rows
        finally:
            await browser.close()


async def run_scrape(
    query: str,
    headless: bool,
    max_results: int,
    delay: float,
    export_format: str,
    proxy: str,
    progress,
    stop: asyncio.Event,
) -> None:
    if scrape_lock.locked():
        await progress(
            {
                "type": "error",
                "message": "Another search is already running. Wait for it to finish.",
            }
        )
        return
    await scrape_lock.acquire()
    try:
        rows = await scrape_maps(query, headless, max_results, delay, proxy, progress, stop)
        if stop.is_set() and not rows:
            return
        if not rows:
            await progress(
                {
                    "type": "error",
                    "message": "No businesses found. Include a city in the search, for example “Konya restoran”.",
                }
            )
            return
        saved = save_export(rows, query, export_format)
        log.info("Saved %s rows to %s", len(rows), saved["file_id"])
        await progress(
            {
                "type": "done",
                "count": len(rows),
                "download_url": saved["download_url"],
                "filename": saved["filename"],
                "cancelled": stop.is_set(),
                "message": f"Saved {len(rows)} businesses to {saved['filename']}",
                "log": True,
            }
        )
    except asyncio.CancelledError:
        raise
    except Exception as exc:
        log.exception("Scrape failed")
        await progress({"type": "error", "message": friendly_error(exc)})
    finally:
        scrape_lock.release()


def file_response(saved: dict[str, str]) -> FileResponse:
    path = resolve_export(saved["file_id"])
    media = (
        "text/csv; charset=utf-8"
        if path.suffix == ".csv"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return FileResponse(path, media_type=media, filename=saved["filename"])


async def scrape_to_file(
    query: str,
    headless: bool,
    max_results: int,
    delay: float,
    export_format: str,
    proxy: str,
) -> FileResponse:
    if scrape_lock.locked():
        raise HTTPException(status_code=409, detail="Another search is already running.")
    await scrape_lock.acquire()
    try:
        rows = await scrape_maps(
            query,
            headless,
            max_results,
            delay,
            proxy,
            progress=noop_progress,
            stop=asyncio.Event(),
        )
    except Exception as exc:
        raise HTTPException(status_code=502, detail=friendly_error(exc)) from exc
    finally:
        scrape_lock.release()
    if not rows:
        raise HTTPException(status_code=404, detail="No businesses found for that search.")
    return file_response(save_export(rows, query, export_format))


async def noop_progress(_: dict) -> None:
    return None


@app.get("/")
async def home() -> FileResponse:
    return FileResponse(TEMPLATE_PATH, headers={"Cache-Control": "no-cache"})


@app.get("/api/search")
async def search_get(
    query: str = Query(..., description="Google Maps search, for example konya restoran"),
    headless: bool = Query(True),
    max_results: int = Query(15, ge=1, le=40),
    delay: float = Query(3, ge=0, le=60),
    export_format: str = Query("xlsx"),
    proxy: str = Query(""),
) -> FileResponse:
    query, delay, export_format, proxy = validate_params(query, max_results, delay, export_format, proxy)
    return await scrape_to_file(query, headless, max_results, delay, export_format, proxy)


@app.post("/api/search")
async def search_post(body: SearchRequest) -> FileResponse:
    query, delay, export_format, proxy = validate_params(
        body.query, body.max_results, body.delay, body.export_format, body.proxy
    )
    return await scrape_to_file(query, body.headless, body.max_results, delay, export_format, proxy)


@app.get("/api/stream")
async def search_stream(
    request: Request,
    query: str = Query(...),
    headless: bool = Query(True),
    max_results: int = Query(15, ge=1, le=40),
    delay: float = Query(3, ge=0, le=60),
    export_format: str = Query("xlsx"),
    proxy: str = Query(""),
) -> StreamingResponse:
    query, delay, export_format, proxy = validate_params(query, max_results, delay, export_format, proxy)
    if scrape_lock.locked():
        raise HTTPException(status_code=409, detail="Another search is already running.")

    queue: asyncio.Queue = asyncio.Queue()
    stop = asyncio.Event()

    async def progress(event: dict) -> None:
        await queue.put(event)

    async def runner() -> None:
        try:
            await run_scrape(query, headless, max_results, delay, export_format, proxy, progress, stop)
        finally:
            await queue.put(None)

    task = asyncio.create_task(runner())

    async def generate():
        try:
            while True:
                if await request.is_disconnected():
                    stop.set()
                    break
                try:
                    item = await asyncio.wait_for(queue.get(), timeout=0.5)
                except asyncio.TimeoutError:
                    yield ": keepalive\n\n"
                    continue
                if item is None:
                    break
                yield f"data: {json.dumps(item, ensure_ascii=False)}\n\n"
        finally:
            stop.set()
            if not task.done():
                try:
                    await asyncio.wait_for(task, timeout=30)
                except asyncio.TimeoutError:
                    task.cancel()

    return StreamingResponse(
        generate(),
        media_type="text/event-stream",
        headers={
            "Cache-Control": "no-cache",
            "Connection": "keep-alive",
            "X-Accel-Buffering": "no",
        },
    )


@app.get("/api/download/{file_id}")
async def download(file_id: str) -> FileResponse:
    path = resolve_export(file_id)
    media = (
        "text/csv; charset=utf-8"
        if path.suffix == ".csv"
        else "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    return FileResponse(path, media_type=media, filename=file_id.split("__", 1)[1])


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("main:app", host="127.0.0.1", port=8000, reload=False)
